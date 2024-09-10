from pptx import Presentation
import argparse
import os

def count_slides(pres) -> int:
    return len(pres.slides)

def grab_title_if_exists(slide) -> str:
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text:
            # Check if the shape is the title
            if shape == slide.shapes.title:
                return shape.text_frame.text
    return ""

def iterate_thru_slides(pres, *funcs):
    # Applies each func passed to each slide, saves as a matrix
    # Rows are slides, cols are slide elements
    return [ [func(slide) for func in funcs]
            for slide in pres.slides ]

def data_to_text_file(data, filename):
    with open(filename, 'w') as file:
        for sublist in data:
            line = ' '.join(map(str, sublist))  # Each item to string and join with spaces
            file.write(line + '\n')  # Write each sublist as a line in the file
    print("Data saved to " + str(filename))


def main():
    # Create Parser, add arguments, parse arguments, create funct descrip.
    my_desc = "A function for grabbing certain basic elements from pptx"
    parser = argparse.ArgumentParser(description = my_desc)
    parser.add_argument("prs_filepath", type=str, help="Path to presentation file.")
    args, unknown = parser.parse_known_args()

    # OPTIONAL: 
    default_output_filename = f"{os.path.splitext(os.path.basename(args.prs_filepath))[0]}-slide_titles"
    parser.add_argument('-o', '--output_filename', type=str, \
                        default=default_output_filename, \
                        help='Specify output filename. \
                        Default is original pptx title plus \'-slide_titles.txt\'.')

    args = parser.parse_args()

    # Main Function: Save all titles from the .pptx files as a .txt list.
    # Only slide titles currently, intend on adding more elements in future.

    pres = Presentation(args.prs_filepath)

    # Grab all titles as a list, write to output text file  
    title_list = iterate_thru_slides(pres, grab_title_if_exists)
    data_to_text_file(title_list, f"{args.output_filename}.txt")

# Run main function if running script directly
if __name__ == "__main__":
    main()

